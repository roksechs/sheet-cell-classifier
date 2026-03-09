"""Evaluate the cell classifier on the DECO dataset.

Usage:
    python scripts/eval_deco.py /tmp/deco/completed

DECO label mapping (binary):
    Header, GroupHead → 0 (header)
    Data, Derived     → 1 (data)
    Table, Notes, Other, MetaTitle → skipped
"""
from __future__ import annotations

import sys
import warnings
from pathlib import Path

import joblib
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sklearn.metrics import classification_report

warnings.filterwarnings("ignore")

_LABEL_MAP = {
    "Header": 0,
    "GroupHead": 0,
    "Data": 1,
    "Derived": 1,
}

_N_FEATURES = 25


def _is_numeric(s: str) -> bool:
    if not s:
        return False
    s = s.strip().replace(",", "").replace("$", "").replace("%", "")
    if s.startswith("(") and s.endswith(")"):
        s = s[1:-1]
    if s.startswith("-"):
        s = s[1:]
    return s.replace(".", "", 1).isdigit()


def _parse_range(range_str: str) -> tuple[int, int, int, int] | None:
    """Parse '$A$1:$C$3' → (min_row, max_row, min_col, max_col)."""
    r = range_str.replace("$", "")
    if ":" not in r:
        return None
    start, end = r.split(":")
    import re
    m1 = re.match(r"([A-Za-z]+)(\d+)", start)
    m2 = re.match(r"([A-Za-z]+)(\d+)", end)
    if not m1 or not m2:
        return None
    c1 = column_index_from_string(m1.group(1))
    r1 = int(m1.group(2))
    c2 = column_index_from_string(m2.group(1))
    r2 = int(m2.group(2))
    return min(r1, r2), max(r1, r2), min(c1, c2), max(c1, c2)


_CAP = 20.0


def extract_from_file(path: Path) -> tuple[np.ndarray, np.ndarray]:
    """Extract (X, y) feature matrix and labels from one DECO xlsx file."""
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        return np.empty((0, _N_FEATURES)), np.empty((0,))

    if "Range_Annotations_Data" not in wb.sheetnames:
        return np.empty((0, _N_FEATURES)), np.empty((0,))

    ann_ws = wb["Range_Annotations_Data"]

    # Parse cell-level annotations (skip Table annotations)
    cell_anns: list[tuple[str, int, str]] = []  # (sheet_name, label, range_str)
    for i, row in enumerate(ann_ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        sheet_name, _, label, _ann_name, rng, _parent, *_ = row
        if not sheet_name or not label or not rng:
            continue
        mapped = _LABEL_MAP.get(label)
        if mapped is not None:
            cell_anns.append((str(sheet_name), mapped, str(rng)))

    rows_X: list[list[float]] = []
    rows_y: list[int] = []

    # Cache per-sheet precomputed context (dist_above, dist_left, row_nf, col_nf, bold, italic)
    sheet_context: dict[str, dict | None] = {}

    for sheet_name, label, rng in cell_anns:
        if sheet_name not in wb.sheetnames:
            continue

        if sheet_name not in sheet_context:
            ws = wb[sheet_name]
            min_row = ws.min_row
            max_row = ws.max_row
            min_col = ws.min_column
            max_col = ws.max_column
            if min_row is None or max_row is None:
                sheet_context[sheet_name] = None
                continue

            n_rows = max_row - min_row + 1
            n_cols = max_col - min_col + 1
            sm: list[list[str]] = [[""] * n_cols for _ in range(n_rows)]
            bold_m: list[list[float]] = [[0.0] * n_cols for _ in range(n_rows)]
            italic_m: list[list[float]] = [[0.0] * n_cols for _ in range(n_rows)]

            for row_cells in ws.iter_rows(
                min_row=min_row, max_row=max_row,
                min_col=min_col, max_col=max_col,
            ):
                for cell in row_cells:
                    ri = cell.row - min_row
                    ci = cell.column - min_col
                    val = cell.value
                    sm[ri][ci] = "" if val is None else str(val)
                    try:
                        bold_m[ri][ci] = float(bool(cell.font and cell.font.bold))
                        italic_m[ri][ci] = float(bool(cell.font and cell.font.italic))
                    except Exception:
                        pass

            # Precompute dist_above and dist_left
            da = [[0] * n_cols for _ in range(n_rows)]
            dl = [[0] * n_cols for _ in range(n_rows)]
            for ri in range(n_rows):
                for ci in range(n_cols):
                    da[ri][ci] = 0 if ri == 0 or sm[ri - 1][ci] == "" else da[ri - 1][ci] + 1
                    dl[ri][ci] = 0 if ci == 0 or sm[ri][ci - 1] == "" else dl[ri][ci - 1] + 1

            # Precompute row and col numeric/text fractions
            row_nf, row_tf = [], []
            for ri in range(n_rows):
                ne = [v for v in sm[ri] if v != ""]
                num = sum(1 for v in ne if _is_numeric(v))
                n = len(ne)
                row_nf.append(num / n if n else 0.0)
                row_tf.append((n - num) / n if n else 0.0)
            col_nf, col_tf = [], []
            for ci in range(n_cols):
                col_vals = [sm[ri][ci] for ri in range(n_rows)]
                ne = [v for v in col_vals if v != ""]
                num = sum(1 for v in ne if _is_numeric(v))
                n = len(ne)
                col_nf.append(num / n if n else 0.0)
                col_tf.append((n - num) / n if n else 0.0)

            sheet_context[sheet_name] = {
                "min_row": min_row, "min_col": min_col,
                "n_rows": n_rows, "n_cols": n_cols,
                "sm": sm, "da": da, "dl": dl,
                "row_nf": row_nf, "col_nf": col_nf,
                "row_tf": row_tf, "col_tf": col_tf,
                "bold_m": bold_m, "italic_m": italic_m,
            }

        ctx = sheet_context.get(sheet_name)
        if ctx is None:
            continue

        parsed = _parse_range(rng)
        if parsed is None:
            continue
        min_row_ann, max_row_ann, min_col_ann, max_col_ann = parsed

        for row_num in range(min_row_ann, max_row_ann + 1):
            for col_num in range(min_col_ann, max_col_ann + 1):
                ri = row_num - ctx["min_row"]
                ci = col_num - ctx["min_col"]
                if ri < 0 or ri >= ctx["n_rows"] or ci < 0 or ci >= ctx["n_cols"]:
                    continue

                val_str = ctx["sm"][ri][ci]
                dist_a = min(float(ctx["da"][ri][ci]), _CAP)
                dist_l = min(float(ctx["dl"][ri][ci]), _CAP)
                rnf = ctx["row_nf"][ri]
                cnf = ctx["col_nf"][ci]
                rtf = ctx["row_tf"][ri]
                ctf = ctx["col_tf"][ci]
                bold = ctx["bold_m"][ri][ci]
                italic = ctx["italic_m"][ri][ci]

                is_num = _is_numeric(val_str)

                x = [0.0] * _N_FEATURES
                x[0] = dist_a
                x[1] = dist_l
                x[2] = dist_a / _CAP
                x[3] = dist_l / _CAP
                x[4] = float(dist_a == 0)
                x[5] = float(dist_l == 0)
                x[6] = float(dist_a < 2)
                x[7] = float(len(val_str) == 0)
                x[8] = float(is_num)
                x[9] = float(len(val_str) > 50)
                x[10] = rnf
                x[11] = cnf
                x[12] = bold
                x[13] = italic
                x[23] = rtf
                x[24] = ctf

                rows_X.append(x)
                rows_y.append(label)

    if not rows_X:
        return np.empty((0, _N_FEATURES)), np.empty((0,))

    return np.array(rows_X, dtype=np.float32), np.array(rows_y, dtype=np.int8)


def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} <deco_completed_dir>")
        sys.exit(1)

    deco_dir = Path(sys.argv[1])
    model_path = Path(__file__).resolve().parent.parent / "src" / "sheet_call_tree" / "cell_classifier.joblib"
    clf = joblib.load(model_path)

    xlsx_files = sorted(deco_dir.glob("*.xlsx"))
    print(f"Found {len(xlsx_files)} DECO files in {deco_dir}")

    all_X: list[np.ndarray] = []
    all_y: list[np.ndarray] = []
    skipped = 0

    for path in xlsx_files:
        X, y = extract_from_file(path)
        if X.size == 0:
            skipped += 1
            continue
        all_X.append(X)
        all_y.append(y)

    if not all_X:
        print("No usable data found.")
        sys.exit(1)

    X_all = np.vstack(all_X)
    y_all = np.concatenate(all_y)

    print(f"Skipped: {skipped} files (no annotations)")
    print(f"Total cells: {len(y_all):,}  (header={np.sum(y_all==0):,}, data={np.sum(y_all==1):,})")

    y_pred = clf.predict(X_all)

    print("\n=== DECO Evaluation ===")
    print(classification_report(y_all, y_pred, target_names=["header", "data"]))


if __name__ == "__main__":
    main()
